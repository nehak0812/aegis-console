"""AI risk: MIT AI Risk Repository (taxonomy + live counts from the published database) and AI incident reports."""
import csv
import io
import re
from collections import Counter

from aegis import db, net
from aegis.collectors.rss import ingest_feed
from aegis.registry import Source, collector

MIT_XLSX = "https://docs.google.com/spreadsheets/d/15LeHcpeuZC9txkvcaMoh3sUhkMvdMMry69xxXL46DT0/export?format=xlsx"
AIID_RSS = "https://incidentdatabase.ai/rss.xml"
AIAAIC_CSV = "https://docs.google.com/spreadsheets/d/1Bn55B4xz21-_Rgdr8BBb2lt0n_4rzLGxFADMlVW0PYI/export?format=csv&gid=888071280"

MIT_DOMAINS = [
    ("1", "Discrimination & Toxicity", [("1.1", "Unfair discrimination and misrepresentation"), ("1.2", "Exposure to toxic content"), ("1.3", "Unequal performance across groups")]),
    ("2", "Privacy & Security", [("2.1", "Compromise of privacy by obtaining, leaking or correctly inferring sensitive information"), ("2.2", "AI system security vulnerabilities and attacks")]),
    ("3", "Misinformation", [("3.1", "False or misleading information"), ("3.2", "Pollution of information ecosystem and loss of consensus reality")]),
    ("4", "Malicious Actors & Misuse", [("4.1", "Disinformation, surveillance, and influence at scale"), ("4.2", "Cyberattacks, weapon development or use, and mass harm"), ("4.3", "Fraud, scams, and targeted manipulation")]),
    ("5", "Human-Computer Interaction", [("5.1", "Overreliance and unsafe use"), ("5.2", "Loss of human agency and autonomy")]),
    ("6", "Socioeconomic & Environmental", [("6.1", "Power centralization and unfair distribution of benefits"), ("6.2", "Increased inequality and decline in employment quality"), ("6.3", "Economic and cultural devaluation of human effort"), ("6.4", "Competitive dynamics"), ("6.5", "Governance failure"), ("6.6", "Environmental harm")]),
    ("7", "AI System Safety, Failures, & Limitations", [("7.1", "AI pursuing its own goals in conflict with human goals or values"), ("7.2", "AI possessing dangerous capabilities"), ("7.3", "Lack of capability or robustness"), ("7.4", "Lack of transparency or interpretability"), ("7.5", "AI welfare and rights"), ("7.6", "Multi-agent risks")]),
]
CYBER_SUBDOMAINS = {"2.1", "2.2", "4.1", "4.2", "4.3", "7.3"}

# first match wins — keyword map from the MIT subdomain definitions
MIT_KEYWORDS = [
    ("2.2", r"prompt injection|jailbreak|model (theft|extraction|inversion)|data poisoning|adversarial (example|attack)|ai agent (hijack|exploit)|llm vulnerab|mcp (server )?vulnerab"),
    ("4.3", r"deepfake (scam|fraud|ceo|call)|voice clon|ai-generated phishing|impersonat|romance scam|sextortion|scam"),
    ("4.2", r"ai-(generated|powered) malware|ai-assisted (hack|exploit|attack)|autonomous (hacking|cyberattack)|bioweapon|chemical weapon|cyberattack"),
    ("4.1", r"influence operation|election|disinformation|facial recognition|surveillance|bot network|propaganda"),
    ("2.1", r"(training data|chat logs|conversations) (leak|exposed)|memoriz|\bpii\b|re-identif|privacy|personal data|data (leak|breach)"),
    ("3.1", r"hallucinat|fabricated (citation|case|quote)|false (answer|claim)|defamat|misinformation"),
    ("3.2", r"ai slop|synthetic content|fake reviews|ai-generated (news|articles)"),
    ("1.1", r"bias(ed)? against|discriminat|racial|gender bias|misgender|stereotyp"),
    ("1.2", r"csam|explicit (images|deepfake)|nonconsensual|non-consensual|hate speech|self-harm|violent content|nudify"),
    ("1.3", r"(fails|worse) (on|for) (darker skin|accents|women|minorities)|error rate"),
    ("5.1", r"over-?reliance|chatbot (advice|told)|suicide|mental health|unsafe (medical|legal) advice|companion"),
    ("5.2", r"manipulat(ive|ion) design|addict|dependence on ai"),
    ("6.2", r"layoff|replac(ed|ing) (workers|jobs)|gig workers|algorithmic (firing|management)"),
    ("6.3", r"copyright|artists|authors sue|training on (books|art)|plagiar"),
    ("6.1", r"monopoly|antitrust|market concentration"),
    ("6.4", r"ai race|rushed release|safety (team|testing) (cut|skipped)"),
    ("6.5", r"regulator|fine[d]?\b|lawsuit|banned|ai act|non-?compliance|ftc"),
    ("6.6", r"data cent(er|re) (water|power|emissions|pollution)|energy consumption"),
    ("7.1", r"reward hacking|scheming|deceptive (alignment|behaviou?r)|resist(s|ed) shutdown"),
    ("7.2", r"self-replicat|autonomous replication|dangerous capabilit|uplift"),
    ("7.6", r"multi-agent|agents (colluding|collude)|agent-to-agent"),
    ("7.3", r"self-driving|autonomous vehicle|robotaxi|crash|malfunction|misclassif|wrong (diagnosis|arrest)|failure|error"),
    ("7.4", r"black box|unexplain|opaque algorithm"),
    ("7.5", r"ai (sentience|consciousness|welfare|rights)"),
]
_MK = [(c, re.compile(p, re.I)) for c, p in MIT_KEYWORDS]


def mit_subdomain(text: str) -> str | None:
    for c, rx in _MK:
        if rx.search(text or ""):
            return c
    return None


@collector(Source(
    id="mit_ai_risk", name="MIT AI Risk Repository", category="AI risk", publisher="MIT FutureTech — AI Risk Repository",
    homepage="https://airisk.mit.edu/", url=MIT_XLSX, cadence_min=10080, licence="CC BY 4.0",
    notes="The 7-domain / 24-subdomain Domain Taxonomy and the Causal Taxonomy, with live counts of catalogued risks per subdomain "
          "from the published database (1,700+ risks from 70+ frameworks)."))
def collect_mit() -> int:
    from openpyxl import load_workbook
    raw = net.cached(MIT_XLSX, 24 * 7, binary=True, timeout=120)
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheet = next((wb[n] for n in wb.sheetnames if "database" in n.lower() and "explainer" not in n.lower()), None)
    counts, causal, total = Counter(), {"Entity": Counter(), "Intent": Counter(), "Timing": Counter()}, 0
    if sheet is not None:
        header, idx = None, {}
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() if c is not None else "" for c in row]
            if header is None:
                low = [c.lower() for c in cells]
                if any(c.startswith("domain") for c in low) and any("sub-domain" in c or "subdomain" in c for c in low):
                    header = cells
                    for i, c in enumerate(low):
                        if c.startswith("sub-domain") or c.startswith("subdomain"):
                            idx.setdefault("sub", i)
                        elif c.startswith("domain"):
                            idx.setdefault("dom", i)
                        elif c in ("entity", "intent", "timing"):
                            idx[c.title()] = i
                continue
            sub = cells[idx["sub"]] if "sub" in idx and idx["sub"] < len(cells) else ""
            m = re.match(r"(\d\.\d)", sub)
            if not m:
                continue
            total += 1
            counts[m.group(1)] += 1
            for k in ("Entity", "Intent", "Timing"):
                if k in idx and idx[k] < len(cells):
                    v = re.sub(r"^\d+\s*-\s*", "", cells[idx[k]]).strip()
                    if v:
                        causal[k][v] += 1
    tax = [{"id": d, "name": n, "cyber": any(s in CYBER_SUBDOMAINS for s, _ in subs),
            "subdomains": [{"id": s, "name": sn, "risks": counts.get(s, 0), "cyber": s in CYBER_SUBDOMAINS} for s, sn in subs]}
           for d, n, subs in MIT_DOMAINS]
    db.kv_set("mit_ai_risk", {"domains": tax, "total_risks": total, "causal": {k: dict(v) for k, v in causal.items()},
                              "source": "https://airisk.mit.edu/", "updated": db.now()})
    return total or len(tax)


@collector(Source(
    id="ai_incidents", name="AI incident reports", category="AI risk", publisher="AI Incident Database (Responsible AI Collaborative)",
    homepage="https://incidentdatabase.ai/", url=AIID_RSS, cadence_min=360, licence="AIID: CC BY-SA 4.0",
    feeds=[{"publisher": "AIAAIC Repository", "url": AIAAIC_CSV}],
    notes="New AI incident reports, classified into MIT AI-risk subdomains by transparent keyword rules. AIAAIC repository used for aggregate context."))
def collect_ai_incidents() -> int:
    n = ingest_feed("ai_incidents", AIID_RSS, "AI Incident Database", "AI incidents", "ai_incident", 100)
    db.x("DELETE FROM item WHERE source_id='ai_incidents' AND (title LIKE 'No title%' OR title='')")
    # classify into MIT subdomains
    rows = db.q("SELECT id, title, summary, entities FROM item WHERE kind='ai_incident'")
    c = db.conn()
    for r in rows:
        ent = r.get("entities") or {}
        ent["mit"] = mit_subdomain(f"{r['title']} {r['summary']}")
        c.execute("UPDATE item SET entities=? WHERE id=?", (__import__("json").dumps(ent), r["id"]))
    c.commit()
    # AIAAIC aggregate context (sector, technology, issue) — counts only
    try:
        txt = net.cached(AIAAIC_CSV, 24 * 3)
        rdr = list(csv.reader(io.StringIO(txt)))
        hi = next(i for i, r in enumerate(rdr[:8]) if any(c.strip().lower() == "headline" for c in r))
        hdr = [h.strip().lower() for h in rdr[hi]]
        col = lambda *names: next((hdr.index(n) for n in names if n in hdr), None)
        cs, ct, ci, cy = col("sector(s)", "sector"), col("technology(ies)", "technology"), col("issue(s)", "ethical issue(s)", "issue"), col("occurred", "year")
        agg = {"sector": Counter(), "technology": Counter(), "issue": Counter(), "year": Counter()}
        total = 0
        for r in rdr[hi + 1:]:
            if len(r) < len(hdr) // 2:
                continue
            total += 1
            for k, ci_ in (("sector", cs), ("technology", ct), ("issue", ci)):
                if ci_ is not None and ci_ < len(r):
                    for v in re.split(r"[;,]\s*", r[ci_]):
                        if v.strip():
                            agg[k][v.strip()] += 1
            if cy is not None and cy < len(r):
                m = re.search(r"(20\d\d)", r[cy])
                if m:
                    agg["year"][m.group(1)] += 1
        db.kv_set("aiaaic", {"total": total, **{k: dict(v.most_common(15)) for k, v in agg.items() if k != "year"},
                             "year": dict(sorted(agg["year"].items())), "source": "https://www.aiaaic.org/aiaaic-repository"})
    except Exception as e:
        print("[aiaaic]", e)
    return n
